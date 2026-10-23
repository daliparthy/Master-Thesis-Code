# python q:\code\hole_analysis.py
r"""
Hole size / hole area analysis for BINARY SEM images (no thresholding).
White = film, black = holes.

Folder layout it expects (same as your SEM data):

    ROOT_DIR\                     e.g. Q:\Gemini SEM\Master Arbeit
        2nm\                      <- thickness folder (number + nm/um)
            vac\                  <- atmosphere folder (the folder right after the thickness folder)
                PlainImages\seg20k\*.tif   (any depth below is fine)
            air\
                ...
        5nm\
            vac\ ...

Every folder below ROOT_DIR that contains images is analysed, image by image.
Results:  ROOT_DIR\Hole_Analysis_Results.xlsx
          sheets: Summary (per thickness + atmosphere), Per image, Per hole
Overlays: a "Hole_Overlays" folder next to each set of images (holes outlined in red)
"""

import os
import re
import glob

import cv2
import numpy as np
import pandas as pd

# =============================================================================
# CONFIGURATION
# =============================================================================

ROOT_DIR = r"Q:\Gemini SEM\Master Arbeit\10nm_depRateOff\10 nm hole sizes"

pixel_size_nm = 5.582

# Black regions smaller than this (in pixels) are ignored as noise.
min_hole_size = 1

# Holes touching the image edge are cut off, so their size is wrong. Skip them?
exclude_edge_holes = True

# True: holes are black on a white film. False: holes are white on black.
holes_are_black = True

# Hole density is given as "holes per this many µm²". 1.0 -> holes/µm²,
# 100.0 -> holes per 100 µm² (10 x 10 µm), etc. Same reference area for every atmosphere.
density_reference_area_um2 = 1.0

save_overlays = True

# Overlay look (colours are B, G, R)
outline_color = (0, 0, 255)       # red outline around counted holes
fill_color = (0, 0, 255)          # tint inside counted holes
fill_alpha = 0.45                 # 0 = no tint, 1 = solid colour
edge_color = (255, 128, 0)        # blue outline around edge holes that were skipped
label_holes = True                # write the Hole ID next to each hole
label_color = (0, 0, 180)
max_labels = 500                  # don't write numbers if an image has more holes than this

image_extensions = (".tif", ".tiff", ".png", ".bmp")

# Folders/files the script should never read as input
skip_folder_names = {"Hole_Overlays"}
skip_file_prefixes = ("overlay_",)

excel_name = "Hole_Analysis_Results.xlsx"

DENSITY_COL = ("Hole density (holes/µm²)" if density_reference_area_um2 == 1
               else f"Hole density (holes per {density_reference_area_um2:g} µm²)")

THICKNESS_RE = re.compile(r"^(\d+(?:[.,]\d+)?)\s*(nm|um|µm|mm)$", re.IGNORECASE)


# =============================================================================
# CORE FUNCTIONS
# =============================================================================

def to_hole_mask(img):
    """Binary image -> uint8 mask with holes = 255, film = 0. No thresholding,
    just splits the two levels already present in the image."""
    if img.ndim == 3:
        img = img[..., :3].max(axis=2) if img.shape[2] >= 3 else img[..., 0]
    lo, hi = img.min(), img.max()
    if lo == hi:                                    # image is all one colour
        white = np.full(img.shape, hi > 0, bool)
    else:
        white = img > (float(lo) + float(hi)) / 2
    holes = ~white if holes_are_black else white
    return holes.astype(np.uint8) * 255


def measure_holes(hole_mask, px_nm):
    """Each connected black region = one hole. Area is the exact pixel count."""
    h, w = hole_mask.shape
    n, labels, stats, centroids = cv2.connectedComponentsWithStats(hole_mask, connectivity=4)

    holes, edge_skipped = [], 0
    drawing = {"counted": [], "edge": [], "labels": []}
    counted = np.zeros(n, bool)
    for i in range(1, n):
        x, y, bw, bh, area_px = stats[i]
        if area_px < min_hole_size:
            continue

        # outline of this hole only (cropped for speed); RETR_CCOMP also gives
        # the outlines of film islands sitting inside a hole
        crop = (labels[y:y + bh, x:x + bw] == i).astype(np.uint8)
        crop = cv2.copyMakeBorder(crop, 1, 1, 1, 1, cv2.BORDER_CONSTANT, value=0)
        cnts, hier = cv2.findContours(crop, cv2.RETR_CCOMP, cv2.CHAIN_APPROX_NONE)
        cnts = [c + np.array([x - 1, y - 1]) for c in cnts]

        on_edge = x == 0 or y == 0 or x + bw == w or y + bh == h
        if exclude_edge_holes and on_edge:
            edge_skipped += 1
            drawing["edge"].extend(cnts)
            continue

        drawing["counted"].extend(cnts)
        counted[i] = True
        drawing["labels"].append((len(holes) + 1, centroids[i]))
        perimeter_px = sum(cv2.arcLength(c, True) for c in cnts)

        ecd_nm = np.sqrt(4 * area_px / np.pi) * px_nm
        circ = 4 * np.pi * area_px / perimeter_px ** 2 if perimeter_px > 0 else np.nan
        holes.append({
            "Hole ID": len(holes) + 1,           # same number as on the overlay image
            "Area (px)": int(area_px),
            "Area (nm²)": area_px * px_nm ** 2,
            "Equivalent diameter (nm)": ecd_nm,
            "Perimeter (nm)": perimeter_px * px_nm,
            "Circularity": min(circ, 1.0) if not np.isnan(circ) else circ,
            "Centroid x (px)": round(float(centroids[i][0]), 1),
            "Centroid y (px)": round(float(centroids[i][1]), 1),
        })
    drawing["counted_mask"] = counted[labels]
    return holes, edge_skipped, drawing


def thickness_and_atmosphere(img_dir, root_dir):
    """Thickness = first folder named like '2nm' / '2.5 nm' below ROOT_DIR.
    Atmosphere = the folder right after it."""
    parts = os.path.relpath(img_dir, root_dir).split(os.sep)
    parts = [p for p in parts if p not in (".", "")]
    for k, p in enumerate(parts):
        m = THICKNESS_RE.match(p.strip())
        if m:
            value = float(m.group(1).replace(",", "."))
            unit = m.group(2).lower().replace("µ", "u")
            atmosphere = parts[k + 1] if k + 1 < len(parts) else "unknown"
            return f"{value:g} {unit}", value, atmosphere
    # no thickness folder found: fall back to first folder as atmosphere
    return "unknown", np.nan, (parts[0] if parts else os.path.basename(root_dir))


def find_image_folders(root_dir):
    folders = {}
    for dirpath, dirnames, filenames in os.walk(root_dir):
        dirnames[:] = [d for d in dirnames if d not in skip_folder_names]
        imgs = sorted(
            os.path.join(dirpath, f) for f in filenames
            if f.lower().endswith(image_extensions) and not f.startswith(skip_file_prefixes)
        )
        if imgs:
            folders[dirpath] = imgs
    return folders


def save_overlay(img, drawing, out_path):
    """Counted holes: red tint + red outline + number (= 'Hole ID' in Excel).
    Edge holes that were skipped: blue outline."""
    img_8u = cv2.normalize(img, None, 0, 255, cv2.NORM_MINMAX, dtype=cv2.CV_8U)
    if img_8u.ndim == 3:
        img_8u = img_8u[..., :3].max(axis=2)
    base = cv2.cvtColor(img_8u, cv2.COLOR_GRAY2BGR)

    # tint the inside of counted holes so you can see exactly what was measured
    overlay = base.copy()
    overlay[drawing["counted_mask"]] = fill_color
    overlay = cv2.addWeighted(overlay, fill_alpha, base, 1 - fill_alpha, 0)

    line = max(2, round(min(base.shape[:2]) / 400))
    cv2.drawContours(overlay, drawing["edge"], -1, edge_color, line)
    cv2.drawContours(overlay, drawing["counted"], -1, outline_color, line)

    if label_holes and len(drawing["labels"]) <= max_labels:
        scale = max(0.35, min(base.shape[:2]) / 1500)
        t = max(1, round(scale * 2))
        for num, (cx, cy) in drawing["labels"]:
            org = (int(cx) + line + 2, int(cy) - line - 2)
            cv2.putText(overlay, str(num), org, cv2.FONT_HERSHEY_SIMPLEX, scale, (255, 255, 255), t + 2, cv2.LINE_AA)
            cv2.putText(overlay, str(num), org, cv2.FONT_HERSHEY_SIMPLEX, scale, label_color, t, cv2.LINE_AA)
    cv2.imwrite(out_path, overlay)


# =============================================================================
# BATCH PROCESSING
# =============================================================================

def process_all(root_dir, px_nm):
    folders = find_image_folders(root_dir)
    if not folders:
        print(f"No images found below {root_dir}")
        return

    image_rows, hole_rows, group_holes = [], [], {}

    for img_dir, image_paths in folders.items():
        thick_label, thick_val, atmosphere = thickness_and_atmosphere(img_dir, root_dir)
        print(f"\n[{thick_label} | {atmosphere}]  {os.path.relpath(img_dir, root_dir)}  "
              f"({len(image_paths)} images)")

        for img_path in image_paths:
            filename = os.path.basename(img_path)
            img = cv2.imread(img_path, cv2.IMREAD_UNCHANGED)
            if img is None:
                print(f"  Could not read {filename}, skipping.")
                continue

            hole_mask = to_hole_mask(img)
            holes, edge_skipped, drawing = measure_holes(hole_mask, px_nm)

            h, w = hole_mask.shape
            image_area_um2 = (h * px_nm * 1e-3) * (w * px_nm * 1e-3)
            areas = np.array([x["Area (nm²)"] for x in holes])
            ecds = np.array([x["Equivalent diameter (nm)"] for x in holes])
            circs = np.array([x["Circularity"] for x in holes])
            count = len(holes)
            # for density every hole counts, also the ones cut by the edge
            # (they are only left out of the SIZE statistics)
            density_count = count + edge_skipped

            image_rows.append({
                "Thickness": thick_label,
                "Atmosphere": atmosphere,
                "Folder": os.path.relpath(img_dir, root_dir),
                "Image": filename,
                "Holes measured (size)": count,
                "Edge holes (not sized)": edge_skipped,
                "Holes counted (density)": density_count,
                "Image area (µm²)": image_area_um2,
                DENSITY_COL: density_count / image_area_um2 * density_reference_area_um2,
                "Mean hole area (nm²)": areas.mean() if count else np.nan,
                "Std hole area (nm²)": areas.std() if count else np.nan,
                "Median hole area (nm²)": np.median(areas) if count else np.nan,
                "Mean hole diameter ECD (nm)": ecds.mean() if count else np.nan,
                "Std hole diameter ECD (nm)": ecds.std() if count else np.nan,
                "Mean circularity": np.nanmean(circs) if count else np.nan,
                "Hole area fraction, all black px (%)": 100 * cv2.countNonZero(hole_mask) / hole_mask.size,
                "_thick_val": thick_val,
            })
            for x in holes:
                hole_rows.append({"Thickness": thick_label, "Atmosphere": atmosphere,
                                  "Image": filename, **x})
            group_holes.setdefault((thick_label, atmosphere), []).extend(holes)

            print(f"  {filename:<45} holes: {count:<6} mean ECD: "
                  f"{(ecds.mean() if count else float('nan')):8.2f} nm")

            if save_overlays:
                ov_dir = os.path.join(img_dir, "Hole_Overlays")
                os.makedirs(ov_dir, exist_ok=True)
                save_overlay(img, drawing,
                             os.path.join(ov_dir, f"overlay_{os.path.splitext(filename)[0]}.png"))

    if not image_rows:
        print("No images could be analysed.")
        return

    # ---- summary per thickness + atmosphere (all holes of the group pooled)
    df_img = pd.DataFrame(image_rows)
    summary_rows = []
    for (thick_label, atmosphere), g in df_img.groupby(["Thickness", "Atmosphere"], sort=False):
        holes = group_holes.get((thick_label, atmosphere), [])
        areas = np.array([x["Area (nm²)"] for x in holes])
        ecds = np.array([x["Equivalent diameter (nm)"] for x in holes])
        n = len(holes)
        total_area = g["Image area (µm²)"].sum()
        total_counted = g["Holes counted (density)"].sum()
        summary_rows.append({
            "Thickness": thick_label,
            "Atmosphere": atmosphere,
            "Images": len(g),
            "Total area analysed (µm²)": total_area,
            "Holes counted (density)": total_counted,
            # total holes / total area: images of every size weigh correctly,
            # and more images does NOT mean a higher density
            DENSITY_COL: total_counted / total_area * density_reference_area_um2,
            f"Std {DENSITY_COL[0].lower()}{DENSITY_COL[1:]} between images":
                g[DENSITY_COL].std(ddof=0),
            "Holes measured (size)": n,
            "Mean hole area (nm²)": areas.mean() if n else np.nan,
            "Std hole area (nm²)": areas.std() if n else np.nan,
            "Median hole area (nm²)": np.median(areas) if n else np.nan,
            "Mean hole diameter ECD (nm)": ecds.mean() if n else np.nan,
            "Std hole diameter ECD (nm)": ecds.std() if n else np.nan,
            "Median hole diameter ECD (nm)": np.median(ecds) if n else np.nan,
            "Min hole diameter ECD (nm)": ecds.min() if n else np.nan,
            "Max hole diameter ECD (nm)": ecds.max() if n else np.nan,
            "Min hole area (nm²)": areas.min() if n else np.nan,
            "Max hole area (nm²)": areas.max() if n else np.nan,
            "Mean hole area fraction (%)": g["Hole area fraction, all black px (%)"].mean(),
            "Std hole area fraction (%)": g["Hole area fraction, all black px (%)"].std(ddof=0),
            "_thick_val": g["_thick_val"].iloc[0],
        })

    summary = (pd.DataFrame(summary_rows)
               .sort_values(["_thick_val", "Atmosphere"], na_position="last")
               .drop(columns="_thick_val"))
    per_image = (df_img.sort_values(["_thick_val", "Atmosphere", "Folder", "Image"], na_position="last")
                 .drop(columns="_thick_val"))
    per_hole = pd.DataFrame(hole_rows)
    conclusions = build_conclusions(summary)

    excel_path = os.path.join(root_dir, excel_name)
    try:
        write_excel(excel_path, summary, per_image, per_hole, conclusions)
    except PermissionError:
        excel_path = excel_path.replace(".xlsx", "_new.xlsx")
        write_excel(excel_path, summary, per_image, per_hole, conclusions)
        print("\n(Results file was open in Excel, saved a new copy instead.)")

    print("\n" + "=" * 60)
    print(summary[["Thickness", "Atmosphere", "Images", "Total area analysed (µm²)", DENSITY_COL,
                   "Mean hole area (nm²)", "Mean hole diameter ECD (nm)"]].to_string(index=False))
    print("=" * 60)
    if (summary["Thickness"] == "unknown").any():
        print("Note: some folders have no thickness folder (like '2nm') in their path.")
    print("\nCONCLUSIONS")
    for line in conclusions:
        print(("  " if line and not line.endswith(":") else "") + line)
    print(f"\nSaved Excel results to: {excel_path}")


def fmt(x, digits=3):
    """Readable number: 0.0123 -> 0.0123, 123.456 -> 123, 12345.6 -> 12346."""
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "n/a"
    return f"{x:.{digits}g}" if abs(x) < 10 ** digits else f"{x:.0f}"


def build_conclusions(summary):
    """One sentence per thickness on hole density, one per atmosphere on hole size range."""
    dens_unit = DENSITY_COL[DENSITY_COL.index("(") + 1:-1]
    lines = []
    for thick, g in summary.groupby("Thickness", sort=False):
        lines.append(f"{thick}:")
        g = g.dropna(subset=[DENSITY_COL])

        # --- density: most vs least
        if len(g) == 0:
            lines.append("No hole density could be calculated.")
        elif len(g) == 1:
            r = g.iloc[0]
            lines.append(f"Only one atmosphere ({r['Atmosphere']}) was measured at this thickness; "
                         f"its hole density is {fmt(r[DENSITY_COL])} {dens_unit}.")
        else:
            d = g[DENSITY_COL]
            hi, lo = g.loc[d.idxmax()], g.loc[d.idxmin()]
            if np.isclose(d.max(), d.min()):
                lines.append(f"All atmospheres have the same hole density "
                             f"({fmt(d.max())} {dens_unit}).")
            else:
                ratio = (f"; that is {d.max() / d.min():.1f}x more holes per area in "
                         f"{hi['Atmosphere']} than in {lo['Atmosphere']}" if d.min() > 0 else "")
                lines.append(
                    f"{hi['Atmosphere']} has the highest hole density ({fmt(hi[DENSITY_COL])} {dens_unit}) "
                    f"and {lo['Atmosphere']} the lowest ({fmt(lo[DENSITY_COL])} {dens_unit}){ratio}.")

        # --- size range per atmosphere
        for _, r in g.iterrows():
            if r["Holes measured (size)"] == 0:
                lines.append(f"{r['Atmosphere']}: no holes were measured, so no size range.")
                continue
            lines.append(
                f"In {r['Atmosphere']}, hole diameters (ECD) range from "
                f"{fmt(r['Min hole diameter ECD (nm)'])} to {fmt(r['Max hole diameter ECD (nm)'])} nm "
                f"(mean {fmt(r['Mean hole diameter ECD (nm)'])} nm, median {fmt(r['Median hole diameter ECD (nm)'])} nm), "
                f"and hole areas from {fmt(r['Min hole area (nm²)'])} to {fmt(r['Max hole area (nm²)'])} nm² "
                f"over {int(r['Holes measured (size)'])} holes.")
        lines.append("")
    return lines[:-1] if lines and lines[-1] == "" else lines


def write_excel(path, summary, per_image, per_hole, conclusions):
    from openpyxl.styles import Font, PatternFill, Alignment
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        # first sheet: plain-language conclusions
        pd.DataFrame({"Conclusions": conclusions}).to_excel(xw, sheet_name="Conclusions", index=False)
        ws = xw.sheets["Conclusions"]
        ws.column_dimensions["A"].width = 130
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="305496")
        for row in ws.iter_rows(min_row=2):
            c = row[0]
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, str) and c.value.endswith(":"):
                c.font = Font(bold=True)

        for name, df in [("Summary", summary), ("Per image", per_image), ("Per hole", per_hole)]:
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="305496")
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 40)
                for c in col[1:]:
                    if isinstance(c.value, float):
                        c.number_format = "0.00"


if __name__ == "__main__":
    process_all(ROOT_DIR, pixel_size_nm)